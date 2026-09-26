"""Turn a repository into retrievable chunks.

Each file type gets a structure-aware splitter so a retrieved chunk is a
meaningful unit (one function, one notebook cell, one README section,
one slide) and carries exact line numbers for citation.
"""
from __future__ import annotations

import ast
import json
import os
import re
import subprocess
from dataclasses import dataclass, field, asdict
from pathlib import Path

MAX_CHARS = 1800      # soft cap per chunk; longer units are split by lines
OVERLAP_LINES = 6     # line overlap between split pieces

SKIP_DIRS = {".git", "__pycache__", ".ipynb_checkpoints", "node_modules",
             "exiftool", ".venv", "venv"}
TEXT_EXT = {".py", ".ipynb", ".md", ".txt", ".rst", ".yml", ".yaml",
            ".toml", ".cfg", ".json", ".sh"}
DOC_EXT = {".pdf", ".pptx"}
MAX_FILE_BYTES = 5_000_000

# Credentials committed to a repo must never be sent to an LLM or shown in answers.
_SECRET_PATTERNS = [
    re.compile(r"""((?:api[_-]?key|token|secret|password|passwd)\s*[=:]\s*)(['"])[^'"\s]{6,}\2""", re.I),
    re.compile(r"(sk-[A-Za-z0-9_\-]{16,})"),
    re.compile(r"(AKIA[0-9A-Z]{16})"),
]


def redact(text: str) -> str:
    text = _SECRET_PATTERNS[0].sub(lambda m: f"{m.group(1)}{m.group(2)}<REDACTED>{m.group(2)}", text)
    for pat in _SECRET_PATTERNS[1:]:
        text = pat.sub("<REDACTED>", text)
    return text


@dataclass
class Chunk:
    repo: str
    path: str               # path relative to repo root
    kind: str               # function | class | module | script-main | cell | section | page | slide
    name: str               # symbol / heading / "cell 4" / "page 2"
    text: str
    start_line: int = 0     # 1-based; 0 when not applicable (pdf/pptx)
    end_line: int = 0
    meta: dict = field(default_factory=dict)

    @property
    def cite(self) -> str:
        m = self.meta or {}
        if "cell" in m:
            loc = f" (cell {m['cell']}" + (f", lines {self.start_line}-{self.end_line})" if self.start_line else ")")
        elif "page" in m:
            loc = f" (page {m['page']})"
        elif "slide" in m:
            loc = f" (slide {m['slide']})"
        elif self.start_line:
            loc = f":L{self.start_line}-{self.end_line}"
        else:
            loc = ""
        return f"{self.repo}/{self.path}{loc}"

    def header(self) -> str:
        return f"[{self.kind}] {self.name} — {self.cite}"

    def to_dict(self):
        return asdict(self)


# ---------------------------------------------------------------- helpers
def _split_lines(lines: list[str], first_line: int, base: dict) -> list[Chunk]:
    """Split a long unit into ~MAX_CHARS pieces on line boundaries with overlap."""
    out, buf, buf_start, size = [], [], 0, 0
    i = 0
    while i < len(lines):
        ln = lines[i]
        if buf and size + len(ln) > MAX_CHARS:
            out.append((buf_start, buf))
            back = max(1, len(buf) - OVERLAP_LINES)
            i = buf_start + back
            buf, size = [], 0
            continue
        if not buf:
            buf_start = i
        buf.append(ln)
        size += len(ln) + 1
        i += 1
    if buf:
        out.append((buf_start, buf))
    chunks = []
    for n, (s, piece) in enumerate(out):
        c = Chunk(**base, text="\n".join(piece),
                  start_line=first_line + s, end_line=first_line + s + len(piece) - 1)
        if len(out) > 1:
            c.name = f"{base['name']} (part {n + 1}/{len(out)})"
        chunks.append(c)
    return chunks


def _unit(repo, path, kind, name, lines, first_line, meta=None) -> list[Chunk]:
    base = dict(repo=repo, path=path, kind=kind, name=name, meta=meta or {})
    text = "\n".join(lines).strip()
    if not text:
        return []
    if len(text) <= MAX_CHARS:
        return [Chunk(**base, text=text, start_line=first_line,
                      end_line=first_line + len(lines) - 1)]
    return _split_lines(lines, first_line, base)


# ---------------------------------------------------------------- python
def load_python(repo: str, path: str, src: str) -> list[Chunk]:
    lines = src.splitlines()
    try:
        tree = ast.parse(src)
    except SyntaxError:
        return _unit(repo, path, "module", path, lines, 1)

    chunks, covered = [], set()
    imports = sorted({ast.unparse(n) for n in tree.body if isinstance(n, (ast.Import, ast.ImportFrom))})
    mod_doc = ast.get_docstring(tree) or ""

    for node in tree.body:
        start = (node.decorator_list[0].lineno if getattr(node, "decorator_list", None) else node.lineno)
        end = node.end_lineno
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            kind = "class" if isinstance(node, ast.ClassDef) else "function"
            # pull in comment lines directly above the definition
            while start > 1 and lines[start - 2].strip().startswith("#"):
                start -= 1
            meta = {"docstring": ast.get_docstring(node) or "",
                    "calls": sorted({n.func.id for n in ast.walk(node)
                                     if isinstance(n, ast.Call) and isinstance(n.func, ast.Name)})}
            if kind == "function":
                meta["signature"] = f"def {node.name}({ast.unparse(node.args)})"
            chunks += _unit(repo, path, kind, node.name, lines[start - 1:end], start, meta)
            covered.update(range(start, end + 1))
        elif (isinstance(node, ast.If) and "__name__" in ast.unparse(node.test)):
            chunks += _unit(repo, path, "script-main", f"{Path(path).name} __main__ (entry point / CLI)",
                            lines[start - 1:end], start)
            covered.update(range(start, end + 1))

    # module-level leftovers (imports, constants, top-level script code),
    # kept as contiguous runs so line citations stay exact
    runs, cur = [], []
    for i in range(1, len(lines) + 1):
        if i in covered:
            if cur:
                runs.append(cur)
            cur = []
        else:
            cur.append(i)
    if cur:
        runs.append(cur)
    for run in runs:
        seg = lines[run[0] - 1:run[-1]]
        if sum(len(l.strip()) for l in seg) < 20:
            continue
        chunks += _unit(repo, path, "module", f"{Path(path).name} module-level code (imports/constants)",
                        seg, run[0], {"docstring": mod_doc})
    for c in chunks:
        c.meta.setdefault("imports", imports)
    return chunks


# ---------------------------------------------------------------- notebooks
def load_notebook(repo: str, path: str, src: str) -> list[Chunk]:
    nb = json.loads(src)
    chunks, last_md = [], ""
    for idx, cell in enumerate(nb.get("cells", [])):
        body = "".join(cell.get("source", []))
        if not body.strip():
            continue
        if cell.get("cell_type") == "markdown":
            last_md = body.strip()
            chunks += _unit(repo, path, "cell", f"markdown cell {idx}", body.splitlines(), 0,
                            {"cell": idx})
            continue
        chunks += _notebook_code_cell(repo, path, idx, body, last_md)
        # stdout of executed cells often records real results (metrics, counts)
        outs = []
        for o in cell.get("outputs", []):
            if o.get("output_type") == "stream":
                outs.append("".join(o.get("text", [])))
            elif "text/plain" in o.get("data", {}):
                outs.append("".join(o["data"]["text/plain"]))
        out_lines = "\n".join(outs).strip().splitlines()
        if len(out_lines) > 70:   # long logs: the start (setup) and end (final metrics) matter most
            out_lines = out_lines[:15] + ["... [log truncated] ..."] + out_lines[-55:]
        if sum(len(l) for l in out_lines) > 80:
            chunks += _unit(repo, path, "cell-output", f"output of code cell {idx}",
                            out_lines, 0, {"cell": idx})
    return chunks


_MAGIC = re.compile(r"^\s*[!%]")


def _notebook_code_cell(repo, path, idx, body, last_md) -> list[Chunk]:
    """Parse a code cell like a Python file so a cell holding many functions
    becomes one chunk per function. Shell/magic lines are commented out for
    parsing only; chunk text keeps the original lines."""
    orig = body.splitlines()
    cleaned = "\n".join("# " + l if _MAGIC.match(l) else l for l in orig)
    base_meta = {"cell": idx}
    if last_md:
        base_meta["context"] = last_md[:300]
    try:
        ast.parse(cleaned)
    except SyntaxError:
        return _unit(repo, path, "cell", f"code cell {idx}", orig, 1, base_meta)
    parts = load_python(repo, path, cleaned)
    for c in parts:
        c.text = "\n".join(orig[c.start_line - 1:c.end_line]).strip() or c.text
        c.meta = {**base_meta, **c.meta}
        if c.kind == "module":
            c.kind, c.name = "cell", f"code cell {idx}"
        elif c.kind == "script-main":
            c.kind, c.name = "cell", f"code cell {idx} (__main__ block)"
    return parts


# ---------------------------------------------------------------- markdown / text
def load_markdown(repo: str, path: str, src: str) -> list[Chunk]:
    lines = src.splitlines()
    chunks, sec_start, heading, trail, in_code = [], 0, Path(path).name, [], False
    def flush(end):
        if end > sec_start:
            name = " > ".join(trail) if trail else heading
            chunks.extend(_unit(repo, path, "section", name, lines[sec_start:end], sec_start + 1))
    for i, ln in enumerate(lines):
        if ln.strip().startswith("```"):
            in_code = not in_code
        m = None if in_code else re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            flush(i)
            level = len(m.group(1))
            trail = trail[:level - 1] + [m.group(2).strip()]
            sec_start = i
    flush(len(lines))
    return chunks


def load_text(repo, path, src):
    return _unit(repo, path, "file", Path(path).name, src.splitlines(), 1)


# ---------------------------------------------------------------- pdf / pptx
def load_pdf(repo: str, path: str, full: Path) -> list[Chunk]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return []
    out = []
    for i, page in enumerate(PdfReader(str(full)).pages, 1):
        txt = (page.extract_text() or "").strip()
        if txt:
            out += _unit(repo, path, "page", f"page {i}", txt.splitlines(), 0, {"page": i})
    return out


def load_pptx(repo: str, path: str, full: Path) -> list[Chunk]:
    try:
        from pptx import Presentation
    except ImportError:
        return []
    out = []
    for i, slide in enumerate(Presentation(str(full)).slides, 1):
        parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Speaker notes: " + notes)
        txt = "\n".join(parts).strip()
        if txt:
            out += _unit(repo, path, "slide", f"slide {i}", txt.splitlines(), 0, {"slide": i})
    return out


# ---------------------------------------------------------------- repo walk
def resolve_repo(src: str, workdir: str = "repos") -> Path:
    """Accept a local path or a git URL (cloned shallowly into workdir)."""
    if os.path.isdir(src):
        return Path(src).resolve()
    name = re.sub(r"\.git$", "", src.rstrip("/").split("/")[-1])
    dest = Path(workdir) / name
    if not dest.exists():
        dest.parent.mkdir(parents=True, exist_ok=True)
        subprocess.run(["git", "clone", "--depth", "1", src, str(dest)], check=True)
    else:
        subprocess.run(["git", "-C", str(dest), "pull", "--ff-only"], check=False,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return dest.resolve()


def load_repo(root: Path, include_docs: bool = True) -> tuple[list[Chunk], list[str]]:
    repo = root.name
    chunks, files = [], []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(d for d in dirnames if d not in SKIP_DIRS)
        for fn in sorted(filenames):
            full = Path(dirpath) / fn
            rel = full.relative_to(root).as_posix()
            ext = full.suffix.lower()
            if full.stat().st_size > MAX_FILE_BYTES and ext != ".ipynb":
                continue
            try:
                if ext in DOC_EXT and include_docs:
                    new = load_pdf(repo, rel, full) if ext == ".pdf" else load_pptx(repo, rel, full)
                elif ext in TEXT_EXT:
                    src = full.read_text(encoding="utf-8", errors="replace")
                    if ext == ".py":
                        new = load_python(repo, rel, src)
                    elif ext == ".ipynb":
                        new = load_notebook(repo, rel, src)
                    elif ext in {".md", ".rst", ".txt"}:
                        new = load_markdown(repo, rel, src)
                    else:
                        new = load_text(repo, rel, src)
                else:
                    continue
            except Exception as e:  # one bad file must not break the index
                print(f"  ! skipped {rel}: {e}")
                continue
            for c in new:
                c.text = redact(c.text)
            if new:
                files.append(rel)
            chunks += new
    return chunks, files


def repo_map(chunks: list[Chunk]) -> str:
    """Compact table of contents: files and the symbols each defines.
    Sent with every question so the model can answer 'what does this repo do' style
    questions and knows what exists even when it wasn't retrieved."""
    by_file: dict[tuple, list[str]] = {}
    for c in chunks:
        key = (c.repo, c.path)
        by_file.setdefault(key, [])
        if c.kind in ("function", "class"):
            by_file[key].append(c.meta.get("signature", f"class {c.name.split(' (part')[0]}"))
    lines = []
    for (repo, path), syms in sorted(by_file.items()):
        lines.append(f"- {repo}/{path}" + (": " + "; ".join(dict.fromkeys(syms)) if syms else ""))
    return "\n".join(lines)
